#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import logging
from io import BytesIO
from pptx import Presentation


class RAGFlowPptParser:
    # Fix P2: removed pointless __init__ that only called super().__init__()
    # on a class with no explicit base class.

    def __get_bulleted_text(self, paragraph):
        is_bulleted = bool(paragraph._p.xpath("./a:pPr/a:buChar")) or bool(paragraph._p.xpath("./a:pPr/a:buAutoNum")) or bool(paragraph._p.xpath("./a:pPr/a:buBlip"))
        if is_bulleted:
            return f"{'  '* paragraph.level}.{paragraph.text}"
        else:
            return paragraph.text

    def __extract(self, shape):
        try:
            # First try to get text content
            if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
                text_frame = shape.text_frame
                texts = []
                for paragraph in text_frame.paragraphs:
                    if paragraph.text.strip():
                        texts.append(self.__get_bulleted_text(paragraph))
                return "\n".join(texts)

            # Safely get shape_type
            try:
                shape_type = shape.shape_type
            except NotImplementedError:
                # If shape_type is not available, try to get text content
                if hasattr(shape, 'text'):
                    return shape.text.strip()
                return ""

            # Handle table
            # Fix P1: guard against None cell text and check value, not cell object truthiness
            if shape_type == 19:
                tb = shape.table
                rows = []
                for i in range(1, len(tb.rows)):
                    row_parts = []
                    for j in range(len(tb.columns)):
                        header = tb.cell(0, j).text or ""
                        value = tb.cell(i, j).text or ""
                        if value:
                            row_parts.append(f"{header}: {value}" if header else value)
                    rows.append("; ".join(row_parts))
                return "\n".join(rows)

            # Handle group shape
            if shape_type == 6:
                texts = []
                for p in sorted(shape.shapes, key=lambda x: (x.top // 10, x.left)):
                    t = self.__extract(p)
                    if t:
                        texts.append(t)
                return "\n".join(texts)

            return ""

        except Exception as e:
            logging.error(f"Error processing shape: {str(e)}")
            return ""

    def __call__(self, fnm, from_page, to_page, callback=None):
        # Fix P3: wrap Presentation open with error handling
        try:
            ppt = Presentation(fnm) if isinstance(fnm, str) else Presentation(BytesIO(fnm))
        except Exception as e:
            logging.error(f"Failed to open PowerPoint file: {e}")
            return []

        txts = []
        self.total_page = len(ppt.slides)
        for i, slide in enumerate(ppt.slides):
            if i < from_page:
                continue
            if i >= to_page:
                break
            texts = []
            try:
                def _safe_sort_key(shape):
                    try:
                        top = shape.top if shape.top is not None else 0
                        left = shape.left if shape.left is not None else 0
                        return (top // 10, left)
                    except (AttributeError, TypeError):
                        return (0, 0)

                for shape in sorted(slide.shapes, key=_safe_sort_key):
                    try:
                        txt = self.__extract(shape)
                        if txt:
                            texts.append(txt)
                    except AttributeError as e:
                        logging.warning(f"Skipping shape on slide {i + 1} due to: {e}")
                        continue
            except AttributeError as e:
                logging.warning(f"Error iterating shapes on slide {i + 1}: {e}")

            # Fix P4: invoke callback to report progress (previously accepted but never called)
            if callback:
                callback((i + 1) / len(ppt.slides), f"Slide {i + 1}/{len(ppt.slides)}")

            txts.append("\n".join(texts))

        return txts
